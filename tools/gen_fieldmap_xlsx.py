# -*- coding: utf-8 -*-
"""《字段地图 v0 确认表》Excel 生成器（对院方推进三件套之三，PRD Q2/D3）
====================================================================
从 demo/data/fieldmap.json（75 份院内真实文书提取的 239 个字段标签）生成
院方逐行勾选确认表：每个字段预填「系统四色建议」，院方核对"是否由 AI 填写"
与下拉选项值后回传，即完成 PRD Q2（变被动为主动）的书面确认闭环。

用法（依赖 openpyxl）：
    python3 tools/gen_fieldmap_xlsx.py
输出：docs/对院方推进/03_字段地图v0_确认表.xlsx
"""
import json, os, sys

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX_SKILL_DIR = os.environ.get("XLSX_SKILL_DIR",
                                "/Users/mac/.zcode/cli/plugins/cache/zcode-plugins-official/spreadsheets/0.1.7/skills/xlsx")
for sub in (XLSX_SKILL_DIR, os.path.join(XLSX_SKILL_DIR, "templates")):
    if sub not in sys.path:
        sys.path.insert(0, sub)
from base import (FONT_NAME, HEADER_BOLD, PRIMARY, NEUTRAL_600, NEUTRAL_900,  # noqa: E402
                  setup_sheet, style_header_row, style_data_row,
                  auto_fit_columns, auto_fit_row_heights,
                  font_body, font_caption, fill_data_row)
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font, Alignment  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402
from openpyxl.worksheet.properties import PageSetupProperties  # noqa: E402

def fit_print(ws, landscape=True):
    """fitToWidth 生效前提是 fitToPage 属性（否则打印/导出 PDF 时列被横向拆页）。"""
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

OUT = os.path.join(REPO, "docs", "对院方推进", "03_字段地图v0_确认表.xlsx")

# ---------------- 四色建议（与 server/records.py field_color 同规则，静态副本防 import 副作用） ----------------
_GREEN_SET = set("主诉 主要症状 症状 持续时间 时间单位 现病史 发病情况 诱因 主要症状特点及其发展变化 伴随症状 上呼吸道症状 消化道症状 泌尿系症状 发病以来诊治经过及结果 发病以来一般情况 精神状态 食欲 睡眠 小便情况 大便情况 体重 与本次疾病无紧密关系的其他疾病情况 病史 既往史 疾病史（含外伤） 一般健康状况标志 健康状况 心血管病史 其他病史 肝炎结核病史 手术外伤输血 手术外伤史 过敏史 预防接种史 个人史 居住地 地址 接触史 疫区接触史 特殊地区居住史 有毒物质接触史 生活习惯、烟酒史 吸烟史 饮酒史 冶游史 婚育史 婚姻史 生育史 月经史 月经量 月经颜色 月经相关症状 家族史 家族健康状况 父母 兄弟姐妹 有无遗传倾向疾病 病史陈述者姓名 病史陈述者 查房记录 初步诊断 入院诊断 出院诊断 诊断依据 鉴别诊断 诊疗计划 诊疗经过 出院情况 出院医嘱 手术经过 术前诊断 术中诊断".split())
_PE_SET = set("发育 营养 表情 面容 神志 体位 配合检查 色泽 肝掌蜘蛛痣 全身浅表淋巴结 头颅异常 眼睑水肿 结膜 巩膜 角膜 瞳孔 对光反射 外耳道 乳突 鼻 鼻窦 口唇 口腔粘膜 齿龈 咽部粘膜 扁桃体 颈部 颈 颈动脉 颈静脉 气管 肝颈静脉回流征 甲状腺 甲状腺异常 胸廓 胸骨叩痛 呼吸运动 呼吸规整 肋间隙 语颤 胸膜摩擦感 叩诊 呼吸音 干湿性罗音 心前区隆起 心律 心包摩擦音 腹外形 腹壁静脉曲张 腹部紧张度 压痛反跳痛 包块 肝脏 肠鸣音 直肠肛门 肛门生殖器 脊柱 脊柱畸形 四肢 专科情况 老中青 起病 护理级别 Padua评分".split())
_VITALS_SET = set("体温 脉搏 呼吸 收缩压 舒张压".split())
_BLUE_EXTRA = ("记录时间", "辅助检查结果", "辅助检查")

CAT_ORDER = ["AI填写", "HIS带入", "模板常规", "留空待补"]

def suggest(label, bindings):
    """四色预建议（与四色确认同口径）：签名/空值类留空，绿=对话提炼，蓝=HIS，灰=模板常规。"""
    label = str(label or "")
    binding = next(iter(bindings), "") if bindings else ""
    if "签名" in label:
        return "留空待补", "医生手工签章，不生成"
    if label in _GREEN_SET:
        return "AI填写", "对话提炼，可溯源到问诊转写"
    if label in _VITALS_SET or label in _BLUE_EXTRA or binding == "Patient":
        return "HIS带入", "院内系统数据自动带入"
    if label in _PE_SET or label.startswith("体格检查"):
        return "模板常规", "按临床常规预填，医生逐项核对"
    return "AI填写", "对话提炼/综合归纳"

def main():
    fm = json.load(open(os.path.join(REPO, "demo", "data", "fieldmap.json"), encoding="utf-8"))
    fields = [f for f in fm.get("fields", []) if str(f.get("label") or "").strip()]
    rows = []
    for f in fields:
        label = str(f["label"]).strip()
        cat, note = suggest(label, f.get("bindings") or {})
        bind_str = "，".join(f"{k}×{v}" for k, v in list((f.get("bindings") or {}).items())[:3])
        rows.append({"label": label, "bind": bind_str, "types": "、".join(f.get("docTypes") or []),
                     "dd": "含下拉" if f.get("dropdown") else "", "cat": cat, "note": note})
    rows.sort(key=lambda r: (CAT_ORDER.index(r["cat"]), r["label"]))

    wb = Workbook()

    # ---------------- Sheet1 填报说明 ----------------
    ws = wb.active
    ws.title = "填报说明"
    n_dd = sum(1 for r in rows if r["dd"])
    stats = [(c, sum(1 for r in rows if r["cat"] == c)) for c in CAT_ORDER]
    setup_sheet(ws, title="字段地图 v0 确认表 · 填报说明", last_col=3)
    info = [
        ("目的", "确认「哪些病历字段由 AI 填写、哪些由系统带入/医生补录」，以及受控下拉字段的选项值——这是病历生成正确性的关键输入（对应字段地图定稿）。"),
        ("数据来源", f"贵院提供的 {fm.get('totalDocs', 75)} 份真实文书样例自动提取，共 {len(rows)} 个字段标签（其中受控下拉 {n_dd} 个），已按系统四色建议预分类。"),
        ("怎么填", "仅需填写「字段确认表」两张灰底列：①【是否由AI填写】选 是/否（不填视为同意系统建议）；②含下拉字段的【下拉选项值核对】列，请补充完整选项值或改正。"),
        ("四色建议口径", "；".join(f"{c}（{dict(stats)[c]}项）" for c in CAT_ORDER)),
        ("统计", f"AI填写 {stats[0][1]} 项 · HIS带入 {stats[1][1]} 项 · 模板常规 {stats[2][1]} 项 · 留空待补 {stats[3][1]} 项"),
        ("回传", "完成后请回传至：【项目负责人/邮箱】；如有整类异议可直接在备注列说明。"),
        ("保密", "本表仅含字段结构与绑定路径，不含任何患者数据。"),
    ]
    for col, h in enumerate(("项目", "说明"), 2):
        ws.cell(row=4, column=col, value=h)
    style_header_row(ws, row_num=4, col_start=2, col_end=3)
    for i, (k, v) in enumerate(info):
        r = 5 + i
        ws.cell(row=r, column=2, value=k)
        ws.cell(row=r, column=3, value=v)
        style_data_row(ws, row_num=r, col_start=2, col_end=3, row_index=i)
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 90
    fit_print(ws, landscape=False)
    auto_fit_row_heights(ws, header_row=4, data_start_row=5)

    # ---------------- Sheet2 字段确认表 ----------------
    ws2 = wb.create_sheet("字段确认表")
    headers = ["序号", "字段标签", "绑定路径（出现次数）", "出现文书类型", "受控下拉",
               "系统四色建议", "是否由AI填写（院方勾选）", "下拉选项值核对（院方填写）", "院方备注"]
    last_col = len(headers) + 1
    setup_sheet(ws2, title="字段地图 v0 确认表（请逐行核对两张灰底列后回传）", last_col=last_col)
    for col, h in enumerate(headers, 2):
        ws2.cell(row=4, column=col, value=h)
    style_header_row(ws2, row_num=4, col_start=2, col_end=last_col)
    for i, r in enumerate(rows):
        rn = 5 + i
        vals = [i + 1, r["label"], r["bind"] or "（无绑定·自由文本）", r["types"], r["dd"],
                f'{r["cat"]}——{r["note"]}', "", "", ""]
        for col, v in enumerate(vals, 2):
            ws2.cell(row=rn, column=col, value=v)
        style_data_row(ws2, row_num=rn, col_start=2, col_end=last_col, row_index=i)
        # 院方填写两列浅底提示
        for col in (8, 9):
            c = ws2.cell(row=rn, column=col)
            c.fill = fill_data_row(1)   # 灰底标识待填
            c.font = Font(name=FONT_NAME, size=11, color=NEUTRAL_600)
    last_row = 4 + len(rows)
    dv = DataValidation(type="list", formula1='"是,否"', allow_blank=True,
                        promptTitle="是否由 AI 填写", prompt="选“是”=同意系统生成该字段；选“否”=由 HIS 带入或医生手工填写")
    ws2.add_data_validation(dv)
    dv.add(f"H5:H{last_row}")
    note_row = last_row + 2
    ws2.cell(row=note_row, column=2,
             value=f"来源：{fm.get('generatedAt', '')} 自动提取自 {fm.get('totalDocs', 75)} 份院内文书样例；"
                   f"“绑定路径”为院内模板 XML 的 DataSource 字段。本表不含患者数据。").font = font_caption()
    ws2.freeze_panes = "D5"
    fit_print(ws2, landscape=True)
    ws2.print_title_rows = "4:4"
    auto_fit_columns(ws2, min_width=8, max_width=34, header_row=4, data_start_row=5)
    # 院方填写三列与建议列给足手写空间（auto_fit 按空数据收窄，需显式覆盖）
    for col, w in {"F": 10, "G": 32, "H": 14, "I": 30, "J": 18}.items():
        ws2.column_dimensions[col].width = w
    auto_fit_row_heights(ws2, header_row=4, data_start_row=5)

    wb.properties.creator = "Z.ai"
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print(f"✓ 生成 {os.path.relpath(OUT, REPO)}（{len(rows)} 字段 × 9 列，含下拉 {n_dd}）")

if __name__ == "__main__":
    main()
